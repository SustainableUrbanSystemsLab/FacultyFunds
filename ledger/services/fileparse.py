"""Read uploaded spreadsheets into rows of plain text.

Supported: .xlsx/.xlsm (openpyxl), legacy .xls (xlrd), PeopleSoft "Excel" downloads that
are really HTML or MHTML, and CSV/TSV text. Every cell becomes a string; dates become ISO text.
"""

import csv
import datetime as dt
import email
import io
import logging
import zipfile
from email import policy
from html.parser import HTMLParser

from django.conf import settings

logger = logging.getLogger(__name__)

MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024


class FileParseError(ValueError):
    """The upload could not be read. The message is safe to show to the user."""


def cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        text = repr(round(value, 9))
        return text[:-2] if text.endswith(".0") else text
    if isinstance(value, int):
        return str(value)
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    return "\n".join(" ".join(line.split()) for line in text.split("\n")).strip()


def read_tables(filename: str, data: bytes) -> dict[str, list[list[str]]]:
    """Return {sheet name: rows}. Raises FileParseError with a user-facing message."""
    if len(data) > settings.IMPORT_MAX_BYTES:
        raise FileParseError("The file is larger than 10 MB.")
    if not data.strip():
        raise FileParseError("The file is empty.")
    try:
        sheets = _dispatch(filename, data)
    except FileParseError:
        raise
    except Exception as exc:  # malformed files can raise almost anything inside the parsers
        logger.warning("Could not parse upload %r: %s", filename, exc)
        raise FileParseError("Could not read this file. Try exporting it again as .xlsx or .csv.") from exc

    cleaned = {}
    for name, rows in sheets.items():
        rows = _trim(rows)
        if rows:
            cleaned[name[:120] or "Sheet"] = rows
    if not cleaned:
        raise FileParseError("No data found in the file.")
    if sum(len(rows) for rows in cleaned.values()) > settings.IMPORT_MAX_ROWS:
        raise FileParseError(f"The file has more than {settings.IMPORT_MAX_ROWS:,} rows.")
    return cleaned


def _dispatch(filename: str, data: bytes) -> dict[str, list[list[str]]]:
    head = data[:4096].lstrip().lower()
    if data[:4] == b"PK\x03\x04":
        return _read_xlsx(data)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return _read_xls(data)
    if head.startswith(b"mime-version") or b"multipart/related" in head:
        return _read_mhtml(data)
    if b"<html" in head or b"<table" in head or head.startswith(b"<!doctype html"):
        return _read_html(_decode_text(data))
    stem = filename.rsplit("/", 1)[-1].rsplit(".", 1)[0] or "Sheet1"
    return {stem: _read_csv(data)}


def _trim(rows: list[list[str]]) -> list[list[str]]:
    rows = [list(row) for row in rows]
    for row in rows:
        while row and not row[-1].strip():
            row.pop()
    while rows and not rows[-1]:
        rows.pop()
    while rows and not rows[0]:
        rows.pop(0)
    return rows


def _decode_text(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _read_csv(data: bytes) -> list[list[str]]:
    text = _decode_text(data)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if sample.count("\t") > sample.count(",") else csv.excel
    return [[cell_to_text(cell) for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data: bytes) -> dict[str, list[list[str]]]:
    import openpyxl  # uses defusedxml when installed

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise FileParseError("This Excel file is too large once uncompressed.")
    except zipfile.BadZipFile as exc:
        raise FileParseError("This doesn't look like a valid Excel file.") from exc

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = {}
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                rows.append([cell_to_text(value) for value in row])
                if len(rows) > settings.IMPORT_MAX_ROWS:
                    raise FileParseError(f"The file has more than {settings.IMPORT_MAX_ROWS:,} rows.")
            sheets[worksheet.title] = rows
        return sheets
    finally:
        workbook.close()


def _read_xls(data: bytes) -> dict[str, list[list[str]]]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except xlrd.XLRDError as exc:
        raise FileParseError("Could not read this .xls file (is it password-protected?).") from exc
    try:
        sheets = {}
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            if sheet.nrows > settings.IMPORT_MAX_ROWS:
                raise FileParseError(f"The file has more than {settings.IMPORT_MAX_ROWS:,} rows.")
            rows = []
            for r in range(sheet.nrows):
                out = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                        except (ValueError, OverflowError, xlrd.xldate.XLDateError):
                            pass
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(cell.value)
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        value = None
                    out.append(cell_to_text(value))
                rows.append(out)
            sheets[sheet.name] = rows
        return sheets
    finally:
        book.release_resources()


def _read_mhtml(data: bytes) -> dict[str, list[list[str]]]:
    message = email.message_from_bytes(data, policy=policy.default)
    for part in message.walk():
        if part.get_content_type() == "text/html":
            return _read_html(part.get_content())
    raise FileParseError("No table found in this file.")


class _TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._open_tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._colspan = 1
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip_depth += 1
        elif tag == "table":
            self._open_tables.append([])
        elif tag == "tr" and self._open_tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
            try:
                self._colspan = max(1, min(int(dict(attrs).get("colspan") or 1), 50))
            except ValueError:
                self._colspan = 1
        elif tag == "br" and self._cell is not None:
            self._cell.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(cell_to_text("".join(self._cell)))
            self._row.extend([""] * (self._colspan - 1))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._open_tables:
            self._open_tables[-1].append(self._row)
            self._row = None
        elif tag == "table" and self._open_tables:
            self.tables.append(self._open_tables.pop())

    def handle_data(self, data):
        if not self._skip_depth and self._cell is not None:
            self._cell.append(data)


def _read_html(text: str) -> dict[str, list[list[str]]]:
    parser = _TableParser()
    parser.feed(text)
    parser.close()
    tables = [table for table in parser.tables if table]
    if not tables:
        raise FileParseError("No table found in this file.")
    return {f"Table {i + 1}": table for i, table in enumerate(tables)}
